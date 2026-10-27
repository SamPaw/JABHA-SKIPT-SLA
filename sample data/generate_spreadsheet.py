import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Student dataset
students = [
    {
        "roll_no": "2025CS001",
        "name": "Aarav Sharma",
        "branch": "Computer Science & Engineering",
        "admission_year": 2025,
        "graduation_year": 2029,
        "s1_os": 78, "s1_js": 85, "s1_dsa": 72,
        "s2_os": 82, "s2_js": 88, "s2_dsa": 79
    },
    {
        "roll_no": "2025CS002",
        "name": "Diya Patel",
        "branch": "Information Technology",
        "admission_year": 2025,
        "graduation_year": 2029,
        "s1_os": 65, "s1_js": 74, "s1_dsa": 58,
        "s2_os": 70, "s2_js": 68, "s2_dsa": 62
    },
    {
        "roll_no": "2025CS003",
        "name": "Rohan Verma",
        "branch": "Computer Science & Engineering",
        "admission_year": 2025,
        "graduation_year": 2029,
        "s1_os": 24, "s1_js": 62, "s1_dsa": 55,  # Fail in S1 OS
        "s2_os": 45, "s2_js": 58, "s2_dsa": 60
    },
    {
        "roll_no": "2025CS004",
        "name": "Ananya Iyer",
        "branch": "Artificial Intelligence & DS",
        "admission_year": 2025,
        "graduation_year": 2029,
        "s1_os": 88, "s1_js": 92, "s1_dsa": 90,
        "s2_os": 85, "s2_js": 95, "s2_dsa": 91
    },
    {
        "roll_no": "2025CS005",
        "name": "Ishaan Nair",
        "branch": "Computer Science & Engineering",
        "admission_year": 2025,
        "graduation_year": 2029,
        "s1_os": 22, "s1_js": 28, "s1_dsa": 45,  # Fail in S1 OS, S1 JS
        "s2_os": 35, "s2_js": 40, "s2_dsa": 50
    },
    {
        "roll_no": "2025CS006",
        "name": "Sneha Kulkarni",
        "branch": "Information Technology",
        "admission_year": 2025,
        "graduation_year": 2029,
        "s1_os": 18, "s1_js": 25, "s1_dsa": 22,  # Fail in S1 OS, S1 JS, S1 DSA
        "s2_os": 26, "s2_js": 32, "s2_dsa": 40   # Fail in S2 OS -> Total 4 fails
    },
    {
        "roll_no": "2025CS007",
        "name": "Aditya Mehta",
        "branch": "Computer Science & Engineering",
        "admission_year": 2025,
        "graduation_year": 2029,
        "s1_os": 55, "s1_js": 64, "s1_dsa": 70,
        "s2_os": 60, "s2_js": 58, "s2_dsa": 65
    },
    {
        "roll_no": "2025CS008",
        "name": "Pooja Reddy",
        "branch": "Artificial Intelligence & DS",
        "admission_year": 2025,
        "graduation_year": 2029,
        "s1_os": 52, "s1_js": 48, "s1_dsa": 25,  # Fail in S1 DSA
        "s2_os": 55, "s2_js": 60, "s2_dsa": 28   # Fail in S2 DSA -> Total 2 fails
    },
    {
        "roll_no": "2025CS009",
        "name": "Vikram Singh",
        "branch": "Computer Science & Engineering",
        "admission_year": 2025,
        "graduation_year": 2029,
        "s1_os": 15, "s1_js": 20, "s1_dsa": 18,  # Fail in S1 OS, S1 JS, S1 DSA
        "s2_os": 22, "s2_js": 25, "s2_dsa": 35   # Fail in S2 OS, S2 JS -> Total 5 fails
    },
    {
        "roll_no": "2025CS010",
        "name": "Riya Gupta",
        "branch": "Information Technology",
        "admission_year": 2025,
        "graduation_year": 2029,
        "s1_os": 72, "s1_js": 80, "s1_dsa": 68,
        "s2_os": 75, "s2_js": 82, "s2_dsa": 70
    }
]

# Calculate fields for CSV & validation
processed_students = []
for s in students:
    marks = [s["s1_os"], s["s1_js"], s["s1_dsa"], s["s2_os"], s["s2_js"], s["s2_dsa"]]
    total = sum(marks)
    percentage = round(total / 6.0, 2)
    failed_marks = [m for m in marks if m < 30]
    backlogs = len(failed_marks)
    
    if backlogs == 0:
        atkt_status = "Not Applicable (Passed All)"
        final_result = "PASS"
    elif backlogs <= 2:
        atkt_status = f"YES - Eligible for ATKT ({backlogs} Backlog{'s' if backlogs > 1 else ''})"
        final_result = "PROMOTED WITH ATKT"
    else:
        atkt_status = f"NO - Ineligible ({backlogs} Backlogs > Max 2)"
        final_result = "FAIL (YEAR DOWN)"
        
    s_proc = dict(s)
    s_proc.update({
        "total_marks": total,
        "percentage": percentage,
        "backlogs": backlogs,
        "atkt_status": atkt_status,
        "final_result": final_result
    })
    processed_students.append(s_proc)

# 1. Export CSV
csv_headers = [
    "Roll No.",
    "Name",
    "Branch",
    "Admission Year",
    "Graduation Year",
    "Sem 1 - OS",
    "Sem 1 - JS",
    "Sem 1 - DSA",
    "Sem 2 - OS",
    "Sem 2 - JS",
    "Sem 2 - DSA",
    "Total Marks (Out of 600)",
    "Percentage (%)",
    "Failed Subjects (Backlogs)",
    "ATKT Eligibility",
    "Final Result"
]

csv_file_path = "/home/sam/js_sla/student_records_atkt.csv"
with open(csv_file_path, "w", newline="", encoding="utf-8") as f:
    writer = csv.writer(f)
    writer.writerow(csv_headers)
    for s in processed_students:
        writer.writerow([
            s["roll_no"],
            s["name"],
            s["branch"],
            s["admission_year"],
            s["graduation_year"],
            s["s1_os"],
            s["s1_js"],
            s["s1_dsa"],
            s["s2_os"],
            s["s2_js"],
            s["s2_dsa"],
            s["total_marks"],
            f"{s['percentage']:.2f}%",
            s["backlogs"],
            s["atkt_status"],
            s["final_result"]
        ])

# Also create students_marks.csv for convenient standard naming
with open("/home/sam/js_sla/students_marks.csv", "w", newline="", encoding="utf-8") as f:
    writer = csv.writer(f)
    writer.writerow(csv_headers)
    for s in processed_students:
        writer.writerow([
            s["roll_no"],
            s["name"],
            s["branch"],
            s["admission_year"],
            s["graduation_year"],
            s["s1_os"],
            s["s1_js"],
            s["s1_dsa"],
            s["s2_os"],
            s["s2_js"],
            s["s2_dsa"],
            s["total_marks"],
            f"{s['percentage']:.2f}%",
            s["backlogs"],
            s["atkt_status"],
            s["final_result"]
        ])

# 2. Build Rich Excel (.xlsx) Workbook
wb = openpyxl.Workbook()

# Style definitions
font_title = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
font_subtitle = Font(name="Calibri", size=11, italic=True, color="E0E0E0")
font_section_header = Font(name="Calibri", size=12, bold=True, color="1F497D")
font_col_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
font_sub_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
font_data = Font(name="Calibri", size=11)
font_bold_data = Font(name="Calibri", size=11, bold=True)
font_fail = Font(name="Calibri", size=11, bold=True, color="9C0006")
font_pass = Font(name="Calibri", size=11, bold=True, color="006100")
font_atkt = Font(name="Calibri", size=11, bold=True, color="9C6500")

fill_title = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
fill_header_main = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
fill_header_sem1 = PatternFill(start_color="337AB7", end_color="337AB7", fill_type="solid")
fill_header_sem2 = PatternFill(start_color="2E8B57", end_color="2E8B57", fill_type="solid")
fill_header_summary = PatternFill(start_color="4B6F96", end_color="4B6F96", fill_type="solid")
fill_zebra = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

fill_fail = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")      # Light red
fill_pass = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")      # Light green
fill_atkt = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")      # Light yellow/amber

fill_box_header = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
fill_box_bg = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

thin_border_side = Side(border_style="thin", color="D3D3D3")
thick_border_side = Side(border_style="medium", color="2F5597")
double_border_side = Side(border_style="double", color="2F5597")

cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
header_border = Border(left=thin_border_side, right=thin_border_side, top=thick_border_side, bottom=thick_border_side)
total_top_border = Border(top=thin_border_side, bottom=double_border_side, left=thin_border_side, right=thin_border_side)

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")

# ==========================================
# SHEET 1: Consolidated Master Report
# ==========================================
ws1 = wb.active
ws1.title = "Master Academic & ATKT Record"
ws1.views.sheetView[0].showGridLines = True

# Title Block
ws1.merge_cells("A1:P1")
title_cell = ws1["A1"]
title_cell.value = "DEPARTMENT OF COMPUTER ENGINEERING & TECHNOLOGY"
title_cell.font = font_title
title_cell.fill = fill_title
title_cell.alignment = align_center
ws1.row_dimensions[1].height = 28

ws1.merge_cells("A2:P2")
sub_cell = ws1["A2"]
sub_cell.value = "Consolidated Student Academic Performance & ATKT Report | Batch: 2025 – 2029 (Passing Marks: 30/100)"
sub_cell.font = font_subtitle
sub_cell.fill = fill_title
sub_cell.alignment = align_center
ws1.row_dimensions[2].height = 20

# Blank Row 3
ws1.row_dimensions[3].height = 8

# Two-tier Headers: Row 4 & Row 5
# Group headers on Row 4
ws1.merge_cells("A4:E4")
ws1["A4"] = "STUDENT PARTICULARS"
ws1["A4"].font = font_col_header
ws1["A4"].fill = fill_header_main
ws1["A4"].alignment = align_center

ws1.merge_cells("F4:H4")
ws1["F4"] = "SEMESTER 1 (MAX: 100)"
ws1["F4"].font = font_col_header
ws1["F4"].fill = fill_header_sem1
ws1["F4"].alignment = align_center

ws1.merge_cells("I4:K4")
ws1["I4"] = "SEMESTER 2 (MAX: 100)"
ws1["I4"].font = font_col_header
ws1["I4"].fill = fill_header_sem2
ws1["I4"].alignment = align_center

ws1.merge_cells("L4:P4")
ws1["L4"] = "OVERALL EVALUATION & ATKT STATUS"
ws1["L4"].font = font_col_header
ws1["L4"].fill = fill_header_summary
ws1["L4"].alignment = align_center

for col in range(1, 17):
    c = ws1.cell(row=4, column=col)
    c.border = header_border
    if col <= 5:
        c.fill = fill_header_main
    elif col <= 8:
        c.fill = fill_header_sem1
    elif col <= 11:
        c.fill = fill_header_sem2
    else:
        c.fill = fill_header_summary

ws1.row_dimensions[4].height = 24

# Row 5 Specific Column Headers
col_names = [
    ("A", "Roll No.", align_center),
    ("B", "Student Name", align_left),
    ("C", "Branch", align_left),
    ("D", "Admission Year", align_center),
    ("E", "Graduation Year", align_center),
    ("F", "OS", align_center),
    ("G", "JS", align_center),
    ("H", "DSA", align_center),
    ("I", "OS", align_center),
    ("J", "JS", align_center),
    ("K", "DSA", align_center),
    ("L", "Total Marks (600)", align_right),
    ("M", "Percentage", align_right),
    ("N", "Backlogs Count", align_center),
    ("O", "ATKT Eligibility", align_center),
    ("P", "Final Result", align_center),
]

for col_idx, (col_letter, name, alignment) in enumerate(col_names, start=1):
    cell = ws1.cell(row=5, column=col_idx, value=name)
    cell.font = font_sub_header
    cell.alignment = alignment
    cell.border = header_border
    if col_idx <= 5:
        cell.fill = fill_header_main
    elif col_idx <= 8:
        cell.fill = fill_header_sem1
    elif col_idx <= 11:
        cell.fill = fill_header_sem2
    else:
        cell.fill = fill_header_summary

ws1.row_dimensions[5].height = 24

# Data Rows 6 to 15
start_row = 6
for i, s in enumerate(students):
    row_num = start_row + i
    is_zebra = (i % 2 == 1)
    base_fill = fill_zebra if is_zebra else fill_white

    ws1.cell(row=row_num, column=1, value=s["roll_no"]).alignment = align_center
    ws1.cell(row=row_num, column=2, value=s["name"]).alignment = align_left
    ws1.cell(row=row_num, column=3, value=s["branch"]).alignment = align_left
    ws1.cell(row=row_num, column=4, value=s["admission_year"]).alignment = align_center
    ws1.cell(row=row_num, column=5, value=s["graduation_year"]).alignment = align_center

    # Marks
    s_marks = [s["s1_os"], s["s1_js"], s["s1_dsa"], s["s2_os"], s["s2_js"], s["s2_dsa"]]
    for m_idx, mark in enumerate(s_marks, start=6):
        m_cell = ws1.cell(row=row_num, column=m_idx, value=mark)
        m_cell.alignment = align_center
        if mark < 30:
            m_cell.font = font_fail
            m_cell.fill = fill_fail
        else:
            m_cell.font = font_data
            m_cell.fill = base_fill

    # Formulas
    # Total Marks: SUM(F{row_num}:K{row_num})
    tot_cell = ws1.cell(row=row_num, column=12, value=f"=SUM(F{row_num}:K{row_num})")
    tot_cell.font = font_bold_data
    tot_cell.alignment = align_right

    # Percentage: L{row_num}/600 formatted as %
    pct_cell = ws1.cell(row=row_num, column=13, value=f"=L{row_num}/600")
    pct_cell.number_format = '0.00%'
    pct_cell.font = font_bold_data
    pct_cell.alignment = align_right

    # Backlogs Count: COUNTIF(F{row_num}:K{row_num}, "<30")
    backlog_cell = ws1.cell(row=row_num, column=14, value=f'=COUNTIF(F{row_num}:K{row_num}, "<30")')
    backlog_cell.font = font_bold_data
    backlog_cell.alignment = align_center

    # ATKT Eligibility: IF(N{row}=0, "Not Required (Passed)", IF(N{row}<=2, "YES - Eligible for ATKT", "NO - Ineligible (> 2 Backlogs)"))
    atkt_formula = f'=IF(N{row_num}=0, "Not Applicable (Passed All)", IF(N{row_num}<=2, "YES - Eligible for ATKT", "NO - Ineligible (> 2 Backlogs)"))'
    atkt_cell = ws1.cell(row=row_num, column=15, value=atkt_formula)
    atkt_cell.alignment = align_center

    # Final Result: IF(N{row}=0, "PASS", IF(N{row}<=2, "PROMOTED WITH ATKT", "FAIL (YEAR DOWN)"))
    res_formula = f'=IF(N{row_num}=0, "PASS", IF(N{row_num}<=2, "PROMOTED WITH ATKT", "FAIL (YEAR DOWN)"))'
    res_cell = ws1.cell(row=row_num, column=16, value=res_formula)
    res_cell.alignment = align_center

    # Style evaluation cells according to student pre-calculated result for optimal display
    proc = processed_students[i]
    if proc["final_result"] == "PASS":
        atkt_cell.font = font_pass
        res_cell.font = font_pass
        res_cell.fill = fill_pass
        backlog_cell.font = font_pass
    elif "ATKT" in proc["final_result"]:
        atkt_cell.font = font_atkt
        atkt_cell.fill = fill_atkt
        res_cell.font = font_atkt
        res_cell.fill = fill_atkt
        backlog_cell.font = font_fail
        backlog_cell.fill = fill_fail
    else:
        atkt_cell.font = font_fail
        atkt_cell.fill = fill_fail
        res_cell.font = font_fail
        res_cell.fill = fill_fail
        backlog_cell.font = font_fail
        backlog_cell.fill = fill_fail

    # Set borders and default fills for metadata cells
    for col_idx in [1, 2, 3, 4, 5, 12, 13]:
        c = ws1.cell(row=row_num, column=col_idx)
        c.fill = base_fill
    for col_idx in range(1, 17):
        ws1.cell(row=row_num, column=col_idx).border = cell_border
        
    ws1.row_dimensions[row_num].height = 22

# Class Average Summary Row (Row 16)
summary_row = 16
ws1.merge_cells(f"A{summary_row}:E{summary_row}")
avg_label = ws1.cell(row=summary_row, column=1, value="CLASS AVERAGE")
avg_label.font = font_bold_data
avg_label.alignment = align_center

for col_idx in range(6, 12):
    col_let = get_column_letter(col_idx)
    avg_cell = ws1.cell(row=summary_row, column=col_idx, value=f"=AVERAGE({col_let}6:{col_let}15)")
    avg_cell.font = font_bold_data
    avg_cell.alignment = align_center
    avg_cell.number_format = '0.0'

# Total Average
tot_avg = ws1.cell(row=summary_row, column=12, value="=AVERAGE(L6:L15)")
tot_avg.font = font_bold_data
tot_avg.alignment = align_right
tot_avg.number_format = '0.0'

# Pct Average
pct_avg = ws1.cell(row=summary_row, column=13, value="=AVERAGE(M6:M15)")
pct_avg.font = font_bold_data
pct_avg.alignment = align_right
pct_avg.number_format = '0.00%'

ws1.cell(row=summary_row, column=14, value=f"=AVERAGE(N6:N15)").alignment = align_center
ws1.cell(row=summary_row, column=14).font = font_bold_data
ws1.cell(row=summary_row, column=14).number_format = '0.0'

ws1.cell(row=summary_row, column=15, value="-").alignment = align_center
ws1.cell(row=summary_row, column=16, value="-").alignment = align_center

for col_idx in range(1, 17):
    c = ws1.cell(row=summary_row, column=col_idx)
    c.border = total_top_border
    c.fill = fill_box_header

ws1.row_dimensions[summary_row].height = 24

# Statistics & Criteria Panel below table
# Row 18: Section Title
ws1.cell(row=18, column=2, value="ACADEMIC SUMMARY & ATKT POLICY STATISTICS").font = font_section_header
ws1.cell(row=18, column=10, value="ATKT (ALLOWED TO KEEP TERM) RULES & CRITERIA").font = font_section_header

# Stat table (Left side: Columns B - F)
stats_data = [
    ("Total Students Enrolled", "=COUNTA(A6:A15)", "10"),
    ("Students Passed (0 Backlogs)", '=COUNTIF(P6:P15, "PASS")', "5 (50.0%)"),
    ("Students Eligible for ATKT (1 - 2 Backlogs)", '=COUNTIF(P6:P15, "PROMOTED WITH ATKT")', "3 (30.0%)"),
    ("Students Failed / Year Down (> 2 Backlogs)", '=COUNTIF(P6:P15, "FAIL (YEAR DOWN)")', "2 (20.0%)"),
    ("Highest Marks Obtained", "=MAX(L6:L15)", "541 / 600 (90.17%)"),
    ("Lowest Marks Obtained", "=MIN(L6:L15)", "135 / 600 (22.50%)")
]

for idx, (label, formula, display_val) in enumerate(stats_data, start=19):
    ws1.merge_cells(f"B{idx}:D{idx}")
    lbl_cell = ws1.cell(row=idx, column=2, value=label)
    lbl_cell.font = font_data
    lbl_cell.border = cell_border
    lbl_cell.fill = fill_white
    
    val_cell = ws1.cell(row=idx, column=5, value=formula)
    val_cell.font = font_bold_data
    val_cell.alignment = align_center
    val_cell.border = cell_border
    val_cell.fill = fill_white
    ws1.row_dimensions[idx].height = 20

# Criteria & Legend Table (Right side: Columns J - P)
rules_data = [
    ("Passing Standard", "Minimum 30 marks out of 100 in each individual subject (OS, JS, DSA). Marks below 30 are treated as Fail / Backlog."),
    ("ATKT Eligibility Criteria", "A student is Allowed To Keep Term (ATKT) if they fail in at most 2 subjects across the academic year."),
    ("Year Down / Ineligible", "Students failing in 3 or more subjects exceed the permissible ATKT limit and must repeat the academic term/year."),
    ("Status Legend: PASS", "Student has passed all 6 subjects (0 backlogs). Promoted unconditionally."),
    ("Status Legend: PROMOTED (ATKT)", "Student has 1 or 2 backlogs. Promoted to Year 2 (2026) with condition to clear backlogs in remedial exams."),
    ("Status Legend: FAIL (YEAR DOWN)", "Student has 3+ backlogs. Ineligible for ATKT and not permitted to proceed to Year 2.")
]

for idx, (rule_title, rule_desc) in enumerate(rules_data, start=19):
    ws1.merge_cells(f"J{idx}:K{idx}")
    t_cell = ws1.cell(row=idx, column=10, value=rule_title)
    t_cell.font = font_bold_data
    t_cell.border = cell_border
    t_cell.fill = fill_white
    
    ws1.merge_cells(f"L{idx}:P{idx}")
    d_cell = ws1.cell(row=idx, column=12, value=rule_desc)
    d_cell.font = font_data
    d_cell.border = cell_border
    d_cell.fill = fill_white
    d_cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws1.row_dimensions[idx].height = 20


# ==========================================
# SHEET 2: Semester 1 Details
# ==========================================
ws2 = wb.create_sheet(title="Semester 1 Details")
ws2.views.sheetView[0].showGridLines = True

ws2.merge_cells("A1:K1")
ws2["A1"] = "SEMESTER 1 SUBJECT PERFORMANCE REPORT"
ws2["A1"].font = font_title
ws2["A1"].fill = fill_header_sem1
ws2["A1"].alignment = align_center
ws2.row_dimensions[1].height = 28

sem1_headers = [
    "Roll No.", "Student Name", "Branch", "Admission Year",
    "OS (100)", "JS (100)", "DSA (100)", "Sem 1 Total (300)",
    "Sem 1 Pct (%)", "Sem 1 Backlogs", "Sem 1 Status"
]

ws2.row_dimensions[3].height = 24
for col_idx, h_text in enumerate(sem1_headers, start=1):
    cell = ws2.cell(row=3, column=col_idx, value=h_text)
    cell.font = font_col_header
    cell.fill = fill_header_sem1
    cell.alignment = align_center
    cell.border = header_border

for i, s in enumerate(students):
    row_num = 4 + i
    is_zebra = (i % 2 == 1)
    base_fill = fill_zebra if is_zebra else fill_white

    ws2.cell(row=row_num, column=1, value=s["roll_no"]).alignment = align_center
    ws2.cell(row=row_num, column=2, value=s["name"]).alignment = align_left
    ws2.cell(row=row_num, column=3, value=s["branch"]).alignment = align_left
    ws2.cell(row=row_num, column=4, value=s["admission_year"]).alignment = align_center

    # Marks
    for c_idx, val in enumerate([s["s1_os"], s["s1_js"], s["s1_dsa"]], start=5):
        m_c = ws2.cell(row=row_num, column=c_idx, value=val)
        m_c.alignment = align_center
        if val < 30:
            m_c.font = font_fail
            m_c.fill = fill_fail
        else:
            m_c.font = font_data
            m_c.fill = base_fill

    # Total & Pct
    tot = ws2.cell(row=row_num, column=8, value=f"=SUM(E{row_num}:G{row_num})")
    tot.font = font_bold_data
    tot.alignment = align_right

    pct = ws2.cell(row=row_num, column=9, value=f"=H{row_num}/300")
    pct.font = font_bold_data
    pct.alignment = align_right
    pct.number_format = '0.00%'

    # Backlogs
    b_c = ws2.cell(row=row_num, column=10, value=f'=COUNTIF(E{row_num}:G{row_num}, "<30")')
    b_c.font = font_bold_data
    b_c.alignment = align_center

    # Status
    st_c = ws2.cell(row=row_num, column=11, value=f'=IF(J{row_num}=0, "CLEAR", IF(J{row_num}<=2, "ATKT", "FAIL"))')
    st_c.alignment = align_center

    s1_fails = [m for m in [s["s1_os"], s["s1_js"], s["s1_dsa"]] if m < 30]
    if len(s1_fails) == 0:
        st_c.font = font_pass
        st_c.fill = fill_pass
        b_c.font = font_pass
    elif len(s1_fails) <= 2:
        st_c.font = font_atkt
        st_c.fill = fill_atkt
        b_c.font = font_fail
        b_c.fill = fill_fail
    else:
        st_c.font = font_fail
        st_c.fill = fill_fail
        b_c.font = font_fail
        b_c.fill = fill_fail

    for c_idx in [1, 2, 3, 4, 8, 9]:
        ws2.cell(row=row_num, column=c_idx).fill = base_fill
    for c_idx in range(1, 12):
        ws2.cell(row=row_num, column=c_idx).border = cell_border
    ws2.row_dimensions[row_num].height = 20


# ==========================================
# SHEET 3: Semester 2 Details
# ==========================================
ws3 = wb.create_sheet(title="Semester 2 Details")
ws3.views.sheetView[0].showGridLines = True

ws3.merge_cells("A1:K1")
ws3["A1"] = "SEMESTER 2 SUBJECT PERFORMANCE REPORT"
ws3["A1"].font = font_title
ws3["A1"].fill = fill_header_sem2
ws3["A1"].alignment = align_center
ws3.row_dimensions[1].height = 28

sem2_headers = [
    "Roll No.", "Student Name", "Branch", "Admission Year",
    "OS (100)", "JS (100)", "DSA (100)", "Sem 2 Total (300)",
    "Sem 2 Pct (%)", "Sem 2 Backlogs", "Sem 2 Status"
]

ws3.row_dimensions[3].height = 24
for col_idx, h_text in enumerate(sem2_headers, start=1):
    cell = ws3.cell(row=3, column=col_idx, value=h_text)
    cell.font = font_col_header
    cell.fill = fill_header_sem2
    cell.alignment = align_center
    cell.border = header_border

for i, s in enumerate(students):
    row_num = 4 + i
    is_zebra = (i % 2 == 1)
    base_fill = fill_zebra if is_zebra else fill_white

    ws3.cell(row=row_num, column=1, value=s["roll_no"]).alignment = align_center
    ws3.cell(row=row_num, column=2, value=s["name"]).alignment = align_left
    ws3.cell(row=row_num, column=3, value=s["branch"]).alignment = align_left
    ws3.cell(row=row_num, column=4, value=s["admission_year"]).alignment = align_center

    # Marks
    for c_idx, val in enumerate([s["s2_os"], s["s2_js"], s["s2_dsa"]], start=5):
        m_c = ws3.cell(row=row_num, column=c_idx, value=val)
        m_c.alignment = align_center
        if val < 30:
            m_c.font = font_fail
            m_c.fill = fill_fail
        else:
            m_c.font = font_data
            m_c.fill = base_fill

    # Total & Pct
    tot = ws3.cell(row=row_num, column=8, value=f"=SUM(E{row_num}:G{row_num})")
    tot.font = font_bold_data
    tot.alignment = align_right

    pct = ws3.cell(row=row_num, column=9, value=f"=H{row_num}/300")
    pct.font = font_bold_data
    pct.alignment = align_right
    pct.number_format = '0.00%'

    # Backlogs
    b_c = ws3.cell(row=row_num, column=10, value=f'=COUNTIF(E{row_num}:G{row_num}, "<30")')
    b_c.font = font_bold_data
    b_c.alignment = align_center

    # Status
    st_c = ws3.cell(row=row_num, column=11, value=f'=IF(J{row_num}=0, "CLEAR", IF(J{row_num}<=2, "ATKT", "FAIL"))')
    st_c.alignment = align_center

    s2_fails = [m for m in [s["s2_os"], s["s2_js"], s["s2_dsa"]] if m < 30]
    if len(s2_fails) == 0:
        st_c.font = font_pass
        st_c.fill = fill_pass
        b_c.font = font_pass
    elif len(s2_fails) <= 2:
        st_c.font = font_atkt
        st_c.fill = fill_atkt
        b_c.font = font_fail
        b_c.fill = fill_fail
    else:
        st_c.font = font_fail
        st_c.fill = fill_fail
        b_c.font = font_fail
        b_c.fill = fill_fail

    for c_idx in [1, 2, 3, 4, 8, 9]:
        ws3.cell(row=row_num, column=c_idx).fill = base_fill
    for c_idx in range(1, 12):
        ws3.cell(row=row_num, column=c_idx).border = cell_border
    ws3.row_dimensions[row_num].height = 20


# Auto-fit column widths across all sheets
for sheet in [ws1, ws2, ws3]:
    for col in sheet.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Skip title row and summary merged rows for column width calculation
            if sheet == ws1 and cell.row in [1, 2, 3, 4, 18, 19, 20, 21, 22, 23, 24]:
                continue
            if (sheet == ws2 or sheet == ws3) and cell.row in [1, 2]:
                continue
            val_str = str(cell.value or '')
            if val_str.startswith("="):
                # Approximation for formula text width
                val_str = "1234567890"
            max_len = max(max_len, len(val_str))
        sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

# Specific column width fine tuning for ws1
ws1.column_dimensions["A"].width = 14
ws1.column_dimensions["B"].width = 20
ws1.column_dimensions["C"].width = 30
ws1.column_dimensions["D"].width = 16
ws1.column_dimensions["E"].width = 16
ws1.column_dimensions["F"].width = 11
ws1.column_dimensions["G"].width = 11
ws1.column_dimensions["H"].width = 11
ws1.column_dimensions["I"].width = 11
ws1.column_dimensions["J"].width = 11
ws1.column_dimensions["K"].width = 11
ws1.column_dimensions["L"].width = 18
ws1.column_dimensions["M"].width = 14
ws1.column_dimensions["N"].width = 16
ws1.column_dimensions["O"].width = 28
ws1.column_dimensions["P"].width = 24

# Save Excel files
xlsx_file_path = "/home/sam/js_sla/students_marks.xlsx"
wb.save(xlsx_file_path)
wb.save("/home/sam/js_sla/student_records_atkt.xlsx")

print("Files successfully generated:")
print("1. /home/sam/js_sla/student_records_atkt.csv")
print("2. /home/sam/js_sla/students_marks.csv")
print("3. /home/sam/js_sla/students_marks.xlsx")
print("4. /home/sam/js_sla/student_records_atkt.xlsx")
